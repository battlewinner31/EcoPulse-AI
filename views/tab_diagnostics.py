# views/tab_diagnostics.py
import streamlit as st
import pandas as pd
from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from core.diagnostics import Diagnostic


def render_diagnostics_tab(diagnostics: list[Diagnostic]):
    st.subheader("AI Fault Intelligence & Optimization")

    if not diagnostics:
        st.success("No active faults or early warnings detected. All systems operating within normal parameters.")
        st.info("EcoPulse AI runs continuous rule-based + statistical analysis after every telemetry cycle.")
        return

    # ------------------------------------------------------------------
    # Summary Metrics
    # ------------------------------------------------------------------
    critical = sum(1 for d in diagnostics if d.severity == "CRITICAL")
    high = sum(1 for d in diagnostics if d.severity == "HIGH")
    early = sum(1 for d in diagnostics if d.severity == "EARLY WARNING")
    total_impact_kw = sum(d.energy_impact_kw for d in diagnostics)
    total_savings_kw = sum(d.estimated_savings_kw for d in diagnostics)
    total_savings_day = sum(d.estimated_savings_inr_per_day for d in diagnostics)
    total_savings_month = sum(d.estimated_savings_inr_per_month for d in diagnostics)

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Active Issues", len(diagnostics))
    c2.metric("Critical / High", f"{critical} / {high}")
    c3.metric("Early Warnings", early)
    c4.metric("Current Extra Power", f"{total_impact_kw:.1f} kW")
    c5.metric("Potential Daily Savings", f"₹{total_savings_day:,.0f}")

    st.markdown("---")

    # ------------------------------------------------------------------
    # Visual Charts
    # ------------------------------------------------------------------
    st.subheader("Energy Impact & Savings Opportunity by Equipment")

    chart_data = []
    for d in diagnostics:
        chart_data.append({
            "Equipment": d.equipment,
            "Current Extra Power (kW)": d.energy_impact_kw,
            "Recoverable Savings (kW)": d.estimated_savings_kw
        })

    if chart_data:
        df_chart = pd.DataFrame(chart_data)
        df_chart = df_chart.groupby("Equipment", as_index=False).sum(numeric_only=True)

        col_left, col_right = st.columns(2)
        with col_left:
            st.markdown("**Current Extra Power Draw**")
            st.bar_chart(df_chart.set_index("Equipment")[["Current Extra Power (kW)"]])
        with col_right:
            st.markdown("**Recoverable Savings After Fix**")
            st.bar_chart(df_chart.set_index("Equipment")[["Recoverable Savings (kW)"]])

        st.caption("Left: where energy is currently being wasted. Right: realistic savings once recommended actions are completed.")

    st.markdown("---")

    # ------------------------------------------------------------------
    # Detailed Cards
    # ------------------------------------------------------------------
    for d in diagnostics:
        color = {
            "CRITICAL": "🔴",
            "HIGH": "🟠",
            "MEDIUM": "🟡",
            "EARLY WARNING": "🟣",
            "LOW": "🔵",
            "INFO": "⚪"
        }.get(d.severity, "⚪")

        with st.expander(
            f"{color} **{d.equipment}** — {d.title}  |  Confidence: {d.confidence:.0%}",
            expanded=(d.severity in ["CRITICAL", "HIGH", "EARLY WARNING"])
        ):
            st.markdown(f"**Severity:** `{d.severity}`")
            st.write(d.description)

            st.markdown("**Probable Causes (ranked)**")
            for cause in d.probable_causes:
                st.write(f"- {cause['cause']} *(confidence {cause['confidence']:.0%})*")

            col_a, col_b, col_c = st.columns(3)
            with col_a:
                st.metric("Extra Power Now", f"{d.energy_impact_kw:.1f} kW")
            with col_b:
                st.metric("Recoverable Savings", f"{d.estimated_savings_kw:.1f} kW")
            with col_c:
                st.metric("Est. Monthly Savings", f"₹{d.estimated_savings_inr_per_month:,.0f}")

            st.markdown("**Immediate Recommended Actions**")
            for i, action in enumerate(d.recommended_actions, 1):
                st.write(f"{i}. {action}")

            if d.optimization_suggestions:
                st.markdown("**Optimization & Setpoint Recommendations**")
                for i, sug in enumerate(d.optimization_suggestions, 1):
                    st.write(f"{i}. {sug}")

            st.caption(f"Responsible team → **{d.responsible_team}**")

    # ------------------------------------------------------------------
    # Total Savings Banner
    # ------------------------------------------------------------------
    st.markdown("---")
    st.success(
        f"**Total Potential Savings if all issues are resolved:**  "
        f"**{total_savings_kw:.1f} kW** continuous  →  "
        f"**₹{total_savings_day:,.0f} / day**  →  "
        f"**₹{total_savings_month:,.0f} / month**"
    )

    # ------------------------------------------------------------------
    # Downloadable Word Document Report
    # ------------------------------------------------------------------
    st.subheader("Download Official EcoPulse AI Report")

    docx_bytes = _build_docx_report(
        diagnostics,
        total_savings_kw,
        total_savings_day,
        total_savings_month
    )

    st.download_button(
        label="Download EcoPulse AI Diagnostic & Optimization Report (.docx)",
        data=docx_bytes,
        file_name=f"EcoPulse_AI_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def _build_docx_report(
    diagnostics: list[Diagnostic],
    total_savings_kw: float,
    total_savings_day: float,
    total_savings_month: float
) -> bytes:
    doc = Document()

    # Title
    title = doc.add_heading("EcoPulse AI", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph("Diagnostic & Optimization Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    # Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(f"Total Active Issues: {len(diagnostics)}")
    doc.add_paragraph(f"Total Recoverable Power: {total_savings_kw:.1f} kW")
    doc.add_paragraph(f"Estimated Daily Savings: ₹{total_savings_day:,.0f}")
    doc.add_paragraph(f"Estimated Monthly Savings: ₹{total_savings_month:,.0f}")
    doc.add_paragraph("")

    # Detailed findings
    doc.add_heading("Detailed Findings & Recommendations", level=1)

    for d in diagnostics:
        doc.add_heading(f"{d.equipment} — {d.title}", level=2)
        doc.add_paragraph(f"Severity: {d.severity}   |   Confidence: {d.confidence:.0%}")
        doc.add_paragraph(d.description)

        doc.add_paragraph("Probable Causes:")
        for c in d.probable_causes:
            doc.add_paragraph(f"• {c['cause']} (confidence {c['confidence']:.0%})", style="List Bullet")

        doc.add_paragraph(f"Current Extra Power: {d.energy_impact_kw:.1f} kW")
        doc.add_paragraph(f"Recoverable Savings: {d.estimated_savings_kw:.1f} kW")
        doc.add_paragraph(f"Est. Daily Savings: ₹{d.estimated_savings_inr_per_day:,.0f}")
        doc.add_paragraph(f"Est. Monthly Savings: ₹{d.estimated_savings_inr_per_month:,.0f}")

        doc.add_paragraph("Immediate Actions:")
        for a in d.recommended_actions:
            doc.add_paragraph(f"• {a}", style="List Bullet")

        if d.optimization_suggestions:
            doc.add_paragraph("Optimization Recommendations:")
            for s in d.optimization_suggestions:
                doc.add_paragraph(f"• {s}", style="List Bullet")

        doc.add_paragraph(f"Responsible Team: {d.responsible_team}")
        doc.add_paragraph("")

    # Final savings box
    doc.add_heading("Total Potential Savings", level=1)
    doc.add_paragraph(f"Continuous power recovery: {total_savings_kw:.1f} kW")
    doc.add_paragraph(f"Daily cost reduction: ₹{total_savings_day:,.0f}")
    doc.add_paragraph(f"Monthly cost reduction: ₹{total_savings_month:,.0f}")
    doc.add_paragraph("")
    doc.add_paragraph(
        "All savings estimates assume the recommended actions are completed "
        "and systems return to design operating points."
    )

    # Footer
    doc.add_paragraph("")
    footer = doc.add_paragraph("Generated by EcoPulse AI — Intelligent Virtual Building Management System")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save to bytes
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()